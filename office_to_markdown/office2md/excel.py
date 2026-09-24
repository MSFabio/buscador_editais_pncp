# -*- coding: utf-8 -*-
"""
Módulo de Conversão de Planilhas Excel (.xlsx, .xls) para Markdown.
"""

from pathlib import Path
from typing import Union, List, Any
import datetime
import openpyxl


def _format_cell_value(val: Any) -> str:
    """Formata valores de células do Excel para texto compatível com Markdown."""
    if val is None:
        return ""

    # Datas e Horários
    if isinstance(val, datetime.datetime):
        if val.time() == datetime.time(0, 0, 0):
            return val.strftime("%Y-%m-%d")
        return val.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(val, datetime.date):
        return val.strftime("%Y-%m-%d")

    # Números decimais inteiros (ex: 154.0 -> 154)
    if isinstance(val, float) and val.is_integer():
        return str(int(val))

    if isinstance(val, float):
        # Arredonda se tiver mais de 4 casas decimais para não poluir
        return f"{val:.4f}".rstrip("0").rstrip(".")

    # Strings: limpa quebras de linha e escapa barras verticais |
    s = str(val).strip()
    lines = [line.strip() for line in s.splitlines() if line.strip()]
    s_cleaned = "<br>".join(lines)
    return s_cleaned.replace("|", "\\|")


def _trim_sheet_data(rows: List[List[Any]]) -> List[List[str]]:
    """Remove linhas e colunas vazias nas bordas para economizar tokens e limpar a tabela."""
    if not rows:
        return []

    # Converte tudo para string formatada
    str_rows = [[_format_cell_value(c) for c in row] for row in rows]

    # Remove linhas vazias no final
    while str_rows and not any(str_rows[-1]):
        str_rows.pop()

    # Remove linhas vazias no início
    while str_rows and not any(str_rows[0]):
        str_rows.pop(0)

    if not str_rows:
        return []

    # Determina colunas vazias
    num_cols = max(len(r) for r in str_rows)
    non_empty_cols = []
    for col_idx in range(num_cols):
        col_has_content = any(len(r) > col_idx and r[col_idx] != "" for r in str_rows)
        if col_has_content:
            non_empty_cols.append(col_idx)

    if not non_empty_cols:
        return []

    min_col = min(non_empty_cols)
    max_col = max(non_empty_cols)

    # Recorta apenas o intervalo com dados úteis
    trimmed = []
    for r in str_rows:
        trimmed_row = []
        for col_idx in range(min_col, max_col + 1):
            val = r[col_idx] if col_idx < len(r) else ""
            trimmed_row.append(val)
        trimmed.append(trimmed_row)

    return trimmed


def _table_to_markdown(data: List[List[str]]) -> str:
    """Converte matriz de strings em tabela Markdown."""
    if not data:
        return "*(Planilha vazia ou sem dados legíveis)*\n"

    headers = data[0]
    # Se o cabeçalho tiver células vazias, preenche com Coluna 1, 2, ...
    safe_headers = []
    for idx, h in enumerate(headers):
        safe_headers.append(h if h else f"Col_{idx + 1}")

    md_lines = []
    md_lines.append("| " + " | ".join(safe_headers) + " |")
    md_lines.append("| " + " | ".join([":---"] * len(safe_headers)) + " |")

    # Linhas de conteúdo
    for row in data[1:]:
        while len(row) < len(safe_headers):
            row.append("")
        md_lines.append("| " + " | ".join(row[:len(safe_headers)]) + " |")

    return "\n".join(md_lines) + "\n"


def convert_xlsx(file_path: Path) -> str:
    """Converte arquivo .xlsx usando openpyxl com leitura de fórmulas avaliadas."""
    wb = openpyxl.load_workbook(str(file_path), data_only=True, read_only=True)
    md_sections = []

    md_sections.append(f"# Pasta de Trabalho: {file_path.name}\n")
    md_sections.append(f"**Abas disponíveis:** {', '.join(wb.sheetnames)}\n")
    md_sections.append("---\n")

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        raw_rows = list(ws.iter_rows(values_only=True))
        trimmed_data = _trim_sheet_data(raw_rows)

        md_sections.append(f"## Planilha: {sheet_name}\n")
        if not trimmed_data:
            md_sections.append("*(Aba sem dados)*\n")
        else:
            md_sections.append(f"> **Dimensões:** {len(trimmed_data)} linhas x {len(trimmed_data[0])} colunas\n")
            md_sections.append(_table_to_markdown(trimmed_data))

    wb.close()
    return "\n".join(md_sections).strip() + "\n"


def convert_xls(file_path: Path) -> str:
    """Converte arquivo legado .xls usando pandas e xlrd."""
    import pandas as pd
    excel_file = pd.ExcelFile(str(file_path), engine="xlrd")
    md_sections = []

    md_sections.append(f"# Pasta de Trabalho (Legado): {file_path.name}\n")
    md_sections.append(f"**Abas disponíveis:** {', '.join(excel_file.sheet_names)}\n")
    md_sections.append("---\n")

    for sheet_name in excel_file.sheet_names:
        df = pd.read_excel(excel_file, sheet_name=sheet_name)
        md_sections.append(f"## Planilha: {sheet_name}\n")

        if df.empty:
            md_sections.append("*(Aba sem dados)*\n")
        else:
            md_sections.append(f"> **Dimensões:** {len(df)} linhas x {len(df.columns)} colunas\n")
            # Converte DataFrame para Markdown
            md_sections.append(df.to_markdown(index=False) + "\n")

    return "\n".join(md_sections).strip() + "\n"


def convert_excel_to_markdown(file_path: Union[str, Path]) -> str:
    """
    Converte uma planilha Excel (.xlsx ou .xls) para Markdown.
    
    Args:
        file_path: Caminho do arquivo .xlsx ou .xls
        
    Returns:
        String contendo a planilha formatada em Markdown
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Arquivo Excel não encontrado: {path}")

    ext = path.suffix.lower()
    if ext == ".xlsx":
        return convert_xlsx(path)
    elif ext == ".xls":
        return convert_xls(path)
    elif ext == ".csv":
        import pandas as pd
        df = pd.read_csv(str(path))
        return f"# Arquivo CSV: {path.name}\n\n" + df.to_markdown(index=False) + "\n"
    else:
        raise ValueError(f"Extensão não suportada para Excel: {ext}. Utilize .xlsx ou .xls")
