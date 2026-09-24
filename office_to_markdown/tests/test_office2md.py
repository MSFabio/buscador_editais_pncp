# -*- coding: utf-8 -*-
"""
Testes Unitários para a biblioteca office2md.
"""

import unittest
from pathlib import Path
import tempfile
import docx
import openpyxl

from office2md import (
    convert_word_to_markdown,
    convert_excel_to_markdown,
    convert_file,
    convert_folder
)


class TestOffice2Md(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.dir_path = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_word_conversion(self):
        # 1. Cria documento Word de teste
        docx_file = self.dir_path / "proposta_comercial.docx"
        doc = docx.Document()
        doc.add_heading("Proposta Comercial K2", level=1)
        doc.add_paragraph("Este documento apresenta os valores de contratação de links SD-WAN.")
        
        p = doc.add_paragraph()
        run = p.add_run("Atenção: ")
        run.bold = True
        run2 = p.add_run("valores sujeitos a alteração.")
        run2.italic = True

        doc.add_paragraph("Primeiro item de benefício", style="List Bullet")
        doc.add_paragraph("Segundo item de benefício", style="List Bullet")

        table = doc.add_table(rows=3, cols=3)
        table.cell(0, 0).text = "Item"
        table.cell(0, 1).text = "Velocidade"
        table.cell(0, 2).text = "Preço Mensal"

        table.cell(1, 0).text = "1"
        table.cell(1, 1).text = "80 Mbps"
        table.cell(1, 2).text = "R$ 653,96"

        table.cell(2, 0).text = "2"
        table.cell(2, 1).text = "100 Mbps"
        table.cell(2, 2).text = "R$ 1.425,37"

        doc.save(str(docx_file))

        # 2. Converte
        res = convert_file(docx_file)
        self.assertEqual(res["status"], "OK")
        self.assertTrue(res["md_path"].exists())

        md_text = res["content"]
        self.assertIn("# Proposta Comercial K2", md_text)
        self.assertIn("**Atenção:**", md_text)
        self.assertIn("*valores sujeitos a alteração.*", md_text)
        self.assertIn("- Primeiro item de benefício", md_text)
        self.assertIn("| Item | Velocidade | Preço Mensal |", md_text)
        self.assertIn("| 1 | 80 Mbps | R$ 653,96 |", md_text)

    def test_excel_conversion(self):
        # 1. Cria planilha Excel de teste
        xlsx_file = self.dir_path / "pesquisa_precos.xlsx"
        wb = openpyxl.Workbook()
        
        # Aba 1
        ws1 = wb.active
        ws1.title = "Links_SDWAN"
        ws1.append(["Lote", "Velocidade", "Qtd", "Valor Unitário", "Valor Total"])
        ws1.append([3, "80 Mbps", 154, 653.965, 100710.61])
        ws1.append([3, "100 Mbps", 7, 1425.3738, 9977.62])

        # Aba 2
        ws2 = wb.create_sheet(title="Resumo_Executivo")
        ws2.append(["Métrica", "Valor"])
        ws2.append(["Total de Links", 161])
        ws2.append(["Custo Mensal Estimado", 110688.23])

        wb.save(str(xlsx_file))

        # 2. Converte
        res = convert_file(xlsx_file)
        self.assertEqual(res["status"], "OK")
        self.assertTrue(res["md_path"].exists())

        md_text = res["content"]
        self.assertIn("## Planilha: Links_SDWAN", md_text)
        self.assertIn("| Lote | Velocidade | Qtd | Valor Unitário | Valor Total |", md_text)
        self.assertIn("80 Mbps", md_text)
        self.assertIn("## Planilha: Resumo_Executivo", md_text)
        self.assertIn("Custo Mensal Estimado", md_text)

    def test_folder_batch_conversion(self):
        # Cria 1 docx e 1 xlsx na pasta
        doc = docx.Document()
        doc.add_heading("Doc 1", level=1)
        doc.save(str(self.dir_path / "doc1.docx"))

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["A", "B"])
        ws.append([1, 2])
        wb.save(str(self.dir_path / "plan1.xlsx"))

        # Converte a pasta inteira
        results = convert_folder(self.dir_path)
        self.assertEqual(len(results), 2)
        for r in results:
            self.assertEqual(r["status"], "OK")
            self.assertTrue(r["md_path"].exists())


if __name__ == "__main__":
    unittest.main()
